from __future__ import annotations
import hashlib
import hmac
import os
import secrets
import sqlite3
import json
import urllib.request
import threading
import time
import base64
import mimetypes
import io
import zipfile
from contextlib import closing
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
def require_admin(request:Request):
    user=current_user(request)
    if not ADMIN_EMAILS or str(user["email"]).lower() not in ADMIN_EMAILS:
        raise HTTPException(403,"Требуются права администратора")
    return user


from fastapi import BackgroundTasks, FastAPI, HTTPException, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse, JSONResponse
from pydantic import BaseModel, Field
from starlette.middleware.sessions import SessionMiddleware

BASE_DIR = Path(__file__).resolve().parent
PROJECT_DIR = BASE_DIR.parent
DATA_DIR = Path(os.getenv("DATA_DIR", PROJECT_DIR / "data"))
UPLOAD_DIR = Path(os.getenv("UPLOAD_DIR", PROJECT_DIR / "uploads"))
FRONTEND_DIR = Path(os.getenv("FRONTEND_DIR", PROJECT_DIR / "frontend"))
DATA_DIR.mkdir(parents=True, exist_ok=True)
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

DB_PATH = DATA_DIR / "chatstudio.db"
SESSION_SECRET = os.getenv("SESSION_SECRET", "").strip()
if len(SESSION_SECRET) < 32:
    raise RuntimeError("SESSION_SECRET must be set to a random value of at least 32 characters")
SESSION_COOKIE_SECURE = os.getenv("SESSION_COOKIE_SECURE", "0").strip().lower() in {"1","true","yes","on"}
CSRF_PROTECTION = os.getenv("CSRF_PROTECTION", "1").strip().lower() in {"1","true","yes","on"}
MAX_NAME_LENGTH = int(os.getenv("MAX_NAME_LENGTH", "80"))
MAX_PROMPT_LENGTH = int(os.getenv("MAX_PROMPT_LENGTH", "30000"))
MAX_SEARCH_LENGTH = int(os.getenv("MAX_SEARCH_LENGTH", "200"))
ADMIN_EMAILS = {x.strip().lower() for x in os.getenv("ADMIN_EMAILS","").split(",") if x.strip()}
MAX_FILE_SIZE = int(os.getenv("MAX_FILE_SIZE", str(20 * 1024 * 1024)))
MAX_TOTAL_FILE_SIZE = int(os.getenv("MAX_TOTAL_FILE_SIZE", str(50 * 1024 * 1024)))
MAX_FILES_PER_REQUEST = int(os.getenv("MAX_FILES_PER_REQUEST", "20"))
ALLOWED_FILE_EXTENSIONS = {".pdf",".docx",".txt",".md",".csv",".xls",".xlsx",".ppt",".pptx",".json",".xml",".zip",".png",".jpg",".jpeg",".gif",".webp"}
GLOBAL_AI_CONCURRENCY = int(os.getenv("GLOBAL_AI_CONCURRENCY", "3"))
REQUEST_TIMEOUT = int(os.getenv("REQUEST_TIMEOUT", "300"))
MAX_FILE_CONTEXT_CHARS = int(os.getenv("MAX_FILE_CONTEXT_CHARS", "500000"))
MAX_ARCHIVE_ENTRIES = int(os.getenv("MAX_ARCHIVE_ENTRIES", "2000"))
MAX_ARCHIVE_UNCOMPRESSED_SIZE = int(os.getenv("MAX_ARCHIVE_UNCOMPRESSED_SIZE", str(100 * 1024 * 1024)))
AI_SEMAPHORE = threading.BoundedSemaphore(max(1, GLOBAL_AI_CONCURRENCY))
AUTH_RATE_WINDOW = int(os.getenv("AUTH_RATE_WINDOW", "900"))
AUTH_LOGIN_LIMIT = int(os.getenv("AUTH_LOGIN_LIMIT", "10"))
AUTH_REGISTER_LIMIT = int(os.getenv("AUTH_REGISTER_LIMIT", "5"))
_AUTH_RATE_LOCK = threading.Lock()
_AUTH_RATE: dict[tuple[str,str], list[float]] = {}

def _client_ip(request: Request) -> str:
    return request.client.host if request.client and request.client.host else "unknown"

def _check_auth_rate(request: Request, action: str, limit: int) -> None:
    now = time.monotonic()
    key = (_client_ip(request), action)
    with _AUTH_RATE_LOCK:
        attempts = [t for t in _AUTH_RATE.get(key, []) if now - t < AUTH_RATE_WINDOW]
        if len(attempts) >= max(1, limit):
            raise HTTPException(429, "Слишком много попыток. Попробуйте позже.")
        attempts.append(now)
        _AUTH_RATE[key] = attempts

def _clear_auth_rate(request: Request, action: str) -> None:
    key = (_client_ip(request), action)
    with _AUTH_RATE_LOCK:
        _AUTH_RATE.pop(key, None)

def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()

def get_db() -> sqlite3.Connection:
    db = sqlite3.connect(DB_PATH)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA foreign_keys = ON")
    return db

def init_db() -> None:
    with closing(get_db()) as db:
        db.executescript("""
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            email TEXT NOT NULL UNIQUE COLLATE NOCASE,
            password_hash TEXT NOT NULL,
            name TEXT NOT NULL,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS chats (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            title TEXT NOT NULL DEFAULT 'Новый чат',
            archived INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS messages (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            chat_id INTEGER NOT NULL,
            user_id INTEGER,
            role TEXT NOT NULL CHECK(role IN ('user','assistant','system')),
            content TEXT NOT NULL,
            model TEXT,
            provider TEXT,
            routing_set TEXT,
            parent_message_id INTEGER,
            created_at TEXT NOT NULL,
            FOREIGN KEY (chat_id) REFERENCES chats(id) ON DELETE CASCADE,
            FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE,
            FOREIGN KEY (parent_message_id) REFERENCES messages(id) ON DELETE SET NULL
        );
        CREATE TABLE IF NOT EXISTS files (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            filename TEXT NOT NULL,
            stored_filename TEXT NOT NULL UNIQUE,
            mime_type TEXT,
            size INTEGER NOT NULL,
            path TEXT NOT NULL,
            created_at TEXT NOT NULL,
            FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS message_files (
            message_id INTEGER NOT NULL,
            file_id INTEGER NOT NULL,
            PRIMARY KEY (message_id,file_id),
            FOREIGN KEY (message_id) REFERENCES messages(id) ON DELETE CASCADE,
            FOREIGN KEY (file_id) REFERENCES files(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS providers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE,
            base_url TEXT NOT NULL,
            api_key_env TEXT,
            enabled INTEGER NOT NULL DEFAULT 1,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS model_configs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            provider_id INTEGER NOT NULL,
            name TEXT NOT NULL,
            model_name TEXT NOT NULL,
            temperature REAL,
            max_tokens INTEGER,
            timeout INTEGER,
            enabled INTEGER NOT NULL DEFAULT 1,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            UNIQUE(provider_id,model_name),
            FOREIGN KEY (provider_id) REFERENCES providers(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS routing_sets (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE,
            description TEXT,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS routing_set_models (
            routing_set_id INTEGER NOT NULL,
            model_config_id INTEGER NOT NULL,
            priority INTEGER NOT NULL,
            PRIMARY KEY (routing_set_id,model_config_id),
            UNIQUE(routing_set_id,priority),
            FOREIGN KEY (routing_set_id) REFERENCES routing_sets(id) ON DELETE CASCADE,
            FOREIGN KEY (model_config_id) REFERENCES model_configs(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS task_routes (
            task_key TEXT PRIMARY KEY,
            routing_set_id INTEGER,
            updated_at TEXT NOT NULL,
            FOREIGN KEY (routing_set_id) REFERENCES routing_sets(id) ON DELETE SET NULL
        );
        CREATE TABLE IF NOT EXISTS chat_shares (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            chat_id INTEGER NOT NULL UNIQUE,
            token TEXT NOT NULL UNIQUE,
            enabled INTEGER NOT NULL DEFAULT 1,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            FOREIGN KEY (chat_id) REFERENCES chats(id) ON DELETE CASCADE
        );
        CREATE INDEX IF NOT EXISTS idx_chat_shares_token ON chat_shares(token);
        CREATE TABLE IF NOT EXISTS chat_suggestions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            chat_id INTEGER NOT NULL,
            suggestions_json TEXT NOT NULL,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            FOREIGN KEY (chat_id) REFERENCES chats(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS ai_requests (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            chat_id INTEGER NOT NULL,
            message_id INTEGER,
            task_key TEXT NOT NULL,
            status TEXT NOT NULL CHECK(status IN ('queued','processing','completed','failed','cancelled')),
            error TEXT,
            provider TEXT,
            model TEXT,
            routing_set TEXT,
            fallback_attempts_json TEXT,
            created_at TEXT NOT NULL,
            started_at TEXT,
            completed_at TEXT,
            cancel_requested INTEGER NOT NULL DEFAULT 0,
            FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE,
            FOREIGN KEY (chat_id) REFERENCES chats(id) ON DELETE CASCADE,
            FOREIGN KEY (message_id) REFERENCES messages(id) ON DELETE SET NULL
        );
        CREATE INDEX IF NOT EXISTS idx_chats_user_updated ON chats(user_id,updated_at DESC);
        CREATE INDEX IF NOT EXISTS idx_messages_chat_created ON messages(chat_id,created_at);
        CREATE INDEX IF NOT EXISTS idx_files_user_created ON files(user_id,created_at DESC);
        CREATE INDEX IF NOT EXISTS idx_ai_requests_user_created ON ai_requests(user_id,created_at DESC);
        CREATE INDEX IF NOT EXISTS idx_ai_requests_chat_created ON ai_requests(chat_id,created_at DESC);
        UPDATE ai_requests SET status='failed',error=COALESCE(error,'Дубликат активного запроса очищен при миграции'),completed_at=COALESCE(completed_at,CURRENT_TIMESTAMP) WHERE task_key='main_generation' AND status IN ('queued','processing') AND id NOT IN (SELECT MAX(id) FROM ai_requests WHERE task_key='main_generation' AND status IN ('queued','processing') GROUP BY chat_id);
        CREATE UNIQUE INDEX IF NOT EXISTS idx_ai_requests_one_active_main_chat ON ai_requests(chat_id) WHERE task_key='main_generation' AND status IN ('queued','processing');
        """)
        columns = {row["name"] for row in db.execute("PRAGMA table_info(ai_requests)").fetchall()}
        if "cancel_requested" not in columns:
            db.execute("ALTER TABLE ai_requests ADD COLUMN cancel_requested INTEGER NOT NULL DEFAULT 0")
        db.commit()

def hash_password(password: str) -> str:
    salt = secrets.token_bytes(16)
    derived = hashlib.scrypt(password.encode(), salt=salt, n=2**14, r=8, p=1)
    return "scrypt$16384$8$1$" + salt.hex() + "$" + derived.hex()

def verify_password(password: str, stored: str) -> bool:
    try:
        algorithm,n,r,p,salt_hex,digest_hex = stored.split("$")
        if algorithm != "scrypt":
            return False
        derived = hashlib.scrypt(password.encode(), salt=bytes.fromhex(salt_hex), n=int(n), r=int(r), p=int(p))
        return hmac.compare_digest(derived.hex(),digest_hex)
    except (ValueError,TypeError):
        return False

def _check_same_origin(request: Request) -> None:
    if not CSRF_PROTECTION:
        return
    origin = request.headers.get("origin")
    if not origin:
        return
    expected = f"{request.url.scheme}://{request.headers.get('host', '')}"
    if origin.rstrip("/") != expected.rstrip("/"):
        raise HTTPException(403, "Недопустимый источник запроса")

def current_user(request: Request) -> sqlite3.Row:
    user_id = request.session.get("user_id")
    if not user_id:
        raise HTTPException(401,"Требуется авторизация")
    with closing(get_db()) as db:
        user = db.execute("SELECT id,email,name,created_at,updated_at FROM users WHERE id=?",(int(user_id),)).fetchone()
    if user is None:
        request.session.clear()
        raise HTTPException(401,"Сессия недействительна")
    return user

def _file_extension(filename: str) -> str:
    return Path(filename).suffix.lower()

def _safe_stored_filename(filename: str) -> str:
    return secrets.token_hex(16) + _file_extension(filename)

def _validate_archive_payload(filename: str, data: bytes) -> None:
    ext=_file_extension(filename)
    if ext not in {".zip",".docx",".xlsx",".pptx"}:
        return
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos=archive.infolist()
            if len(infos)>MAX_ARCHIVE_ENTRIES:
                raise HTTPException(413,"Архив содержит слишком много файлов")
            total=0
            for info in infos:
                if info.flag_bits & 0x1:
                    raise HTTPException(400,"Зашифрованные архивы не поддерживаются")
                total += max(0,int(info.file_size))
                if total>MAX_ARCHIVE_UNCOMPRESSED_SIZE:
                    raise HTTPException(413,"Распакованный размер файла слишком большой")
    except zipfile.BadZipFile:
        raise HTTPException(400,"Файл повреждён или имеет неверный формат")

def _stored_file_path(path_value: str) -> Path:
    base = UPLOAD_DIR.resolve()
    path = Path(path_value)
    if not path.is_absolute():
        path = PROJECT_DIR / path
    resolved = path.resolve()
    try:
        resolved.relative_to(base)
    except ValueError:
        raise HTTPException(404, "Файл не найден")
    return resolved

def _read_file_for_ai(row: sqlite3.Row):
    try:
        path=_stored_file_path(row["path"])
    except HTTPException:
        return None
    if not path.is_file(): return None
    mime=(row["mime_type"] or mimetypes.guess_type(row["filename"])[0] or "application/octet-stream").lower()
    ext=_file_extension(row["filename"])
    if mime.startswith("image/"):
        encoded=base64.b64encode(path.read_bytes()).decode("ascii")
        return {"type":"image_url","image_url":{"url":f"data:{mime};base64,{encoded}"}}
    try:
        if ext in {".txt",".md",".csv",".json",".xml"}:
            text=path.read_text("utf-8",errors="replace")
            return {"type":"text","text":f"Файл {row['filename']}:\n{text[:200000]}"}
        if ext==".pdf":
            from pypdf import PdfReader
            text="\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages)
            return {"type":"text","text":f"Файл {row['filename']} (извлечённый текст):\n{text[:200000]}"}
        if ext==".docx":
            from docx import Document
            text="\n".join(p.text for p in Document(str(path)).paragraphs)
            return {"type":"text","text":f"Файл {row['filename']} (извлечённый текст):\n{text[:200000]}"}
        if ext in {".xls",".xlsx"}:
            from openpyxl import load_workbook
            wb=load_workbook(str(path),read_only=True,data_only=True)
            chunks=[]
            for ws in wb.worksheets:
                chunks.append(f"[Лист: {ws.title}]")
                for values in ws.iter_rows(values_only=True):
                    chunks.append(" | ".join("" if v is None else str(v) for v in values))
                    if len("\n".join(chunks))>200000: break
            return {"type":"text","text":f"Файл {row['filename']}:\n" + "\n".join(chunks)[:200000]}
        if ext in {".ppt",".pptx"}:
            from pptx import Presentation
            chunks=[]
            for slide in Presentation(str(path)).slides:
                for shape in slide.shapes:
                    if hasattr(shape,"text") and shape.text: chunks.append(shape.text)
            return {"type":"text","text":f"Файл {row['filename']} (извлечённый текст):\n" + "\n".join(chunks)[:200000]}
    except Exception as exc:
        return {"type":"text","text":f"Файл {row['filename']} не удалось полностью прочитать: {str(exc)[:300]}"}
    return {"type":"text","text":f"Файл {row['filename']} прикреплён, но его содержимое этого типа пока не извлекается."}

def _chat_ai_messages(db: sqlite3.Connection, chat_id: int):
    rows=db.execute("SELECT id,role,content FROM messages WHERE chat_id=? ORDER BY created_at ASC,id ASC",(chat_id,)).fetchall()
    result=[]
    file_context_used=0
    for row in rows:
        if row["role"]!="user":
            result.append({"role":row["role"],"content":row["content"]})
            continue
        files=db.execute("SELECT f.id,f.filename,f.mime_type,f.size,f.path FROM message_files mf JOIN files f ON f.id=mf.file_id WHERE mf.message_id=? ORDER BY f.id",(row["id"],)).fetchall()
        parts=[]
        if row["content"]: parts.append({"type":"text","text":row["content"]})
        for file_row in files:
            if file_context_used >= MAX_FILE_CONTEXT_CHARS:
                parts.append({"type":"text","text":"[Дополнительные файлы пропущены: достигнут лимит контекста вложений.]"})
                break
            part=_read_file_for_ai(file_row)
            if not part: continue
            if part.get("type")=="text":
                text_value=str(part.get("text",""))
                remaining=max(0,MAX_FILE_CONTEXT_CHARS-file_context_used)
                if len(text_value)>remaining:
                    text_value=text_value[:remaining]+ "\n[Содержимое файла обрезано из-за общего лимита контекста вложений.]"
                file_context_used += len(text_value)
                part={"type":"text","text":text_value}
            elif part.get("type")=="image_url":
                # Image payload size is bounded by the normal upload limit; count a conservative
                # representation budget so a chat with many images cannot grow without bound.
                file_context_used += min(int(file_row["size"]),MAX_FILE_CONTEXT_CHARS-file_context_used)
            parts.append(part)
        result.append({"role":"user","content":parts if len(parts)>1 else (parts[0]["text"] if parts and parts[0]["type"]=="text" else parts)})
    return result

def get_owned_chat(db: sqlite3.Connection, chat_id: int, user_id: int) -> sqlite3.Row:
    chat = db.execute(
        "SELECT id,user_id,title,archived,created_at,updated_at FROM chats WHERE id=? AND user_id=?",
        (chat_id,user_id)
    ).fetchone()
    if chat is None:
        raise HTTPException(404,"Чат не найден")
    return chat

class RegisterRequest(BaseModel):
    name: str = Field(min_length=1,max_length=MAX_NAME_LENGTH)
    email: str = Field(min_length=3,max_length=320)
    password: str = Field(min_length=8,max_length=256)

class LoginRequest(BaseModel):
    email: str
    password: str

class CreateChatRequest(BaseModel):
    title: str = Field(default="Новый чат",min_length=1,max_length=200)

class ShareUpdateRequest(BaseModel):
    enabled: bool = True

class RenameChatRequest(BaseModel):
    title: str = Field(min_length=1,max_length=200)

class ProfileUpdateRequest(BaseModel):
    name: str = Field(min_length=1,max_length=MAX_NAME_LENGTH)

class PasswordUpdateRequest(BaseModel):
    current_password: str
    new_password: str = Field(min_length=8,max_length=256)

class MessageRequest(BaseModel):
    chat_id: int
    content: str = Field(default="",max_length=MAX_PROMPT_LENGTH)
    parent_message_id: int | None = None
    file_ids: list[int] = Field(default_factory=list,max_length=MAX_FILES_PER_REQUEST)
    edit_message_id: int | None = None

class RoutingSetRequest(BaseModel):
    name: str = Field(min_length=1,max_length=120)
    description: str | None = Field(default=None,max_length=500)

class ProviderRequest(BaseModel):
    name: str = Field(min_length=1,max_length=120)
    base_url: str = Field(min_length=1,max_length=500)
    api_key_env: str | None = Field(default=None,max_length=120)

class ModelConfigRequest(BaseModel):
    provider_id: int
    name: str = Field(min_length=1,max_length=120)
    model_name: str = Field(min_length=1,max_length=200)
    temperature: float | None = Field(default=None,ge=0,max=2)
    max_tokens: int | None = Field(default=None,ge=1)
    timeout: int | None = Field(default=None,ge=1)

class ProviderUpdateRequest(BaseModel):
    name: str | None = Field(default=None,min_length=1,max_length=120)
    base_url: str | None = Field(default=None,min_length=1,max_length=500)
    api_key_env: str | None = Field(default=None,max_length=120)
    enabled: bool | None = None

class ModelConfigUpdateRequest(BaseModel):
    name: str | None = Field(default=None,min_length=1,max_length=120)
    model_name: str | None = Field(default=None,min_length=1,max_length=200)
    temperature: float | None = Field(default=None,ge=0,max=2)
    max_tokens: int | None = Field(default=None,ge=1)
    timeout: int | None = Field(default=None,ge=1)
    enabled: bool | None = None

class RoutingSetModelRequest(BaseModel):
    model_config_id: int
    priority: int = Field(ge=1)

class TaskRouteRequest(BaseModel):
    routing_set_id: int | None = None

class ModelRouter:
    def _models_for_task(self, db: sqlite3.Connection, task_key: str):
        row=db.execute("SELECT routing_set_id FROM task_routes WHERE task_key=?",(task_key,)).fetchone()
        if not row or row["routing_set_id"] is None:
            raise HTTPException(503,f"Для задачи {task_key} не настроен набор моделей")
        models=db.execute("""
            SELECT mc.id,mc.name,mc.model_name,mc.temperature,mc.max_tokens,mc.timeout,
                   p.name AS provider_name,p.base_url,p.api_key_env
            FROM routing_set_models rsm
            JOIN model_configs mc ON mc.id=rsm.model_config_id
            JOIN providers p ON p.id=mc.provider_id
            WHERE rsm.routing_set_id=? AND mc.enabled=1 AND p.enabled=1
            ORDER BY rsm.priority ASC
        """,(row["routing_set_id"],)).fetchall()
        return models,int(row["routing_set_id"])

    def _request(self, model, messages):
        api_key=os.getenv(model["api_key_env"]) if model["api_key_env"] else None
        headers={"Content-Type":"application/json"}
        if api_key:
            headers["Authorization"]="Bearer "+api_key
        payload={"model":model["model_name"],"messages":messages}
        if model["temperature"] is not None:
            payload["temperature"]=model["temperature"]
        if model["max_tokens"] is not None:
            payload["max_tokens"]=model["max_tokens"]
        timeout=int(model["timeout"] or REQUEST_TIMEOUT)
        url=model["base_url"].rstrip("/")+"/chat/completions"
        req=urllib.request.Request(url,data=json.dumps(payload).encode("utf-8"),headers=headers,method="POST")
        with urllib.request.urlopen(req,timeout=timeout) as response:
            data=json.loads(response.read().decode("utf-8"))
        content=data.get("choices",[{}])[0].get("message",{}).get("content")
        if not isinstance(content,str) or not content.strip():
            raise RuntimeError("Провайдер вернул пустой ответ")
        return content

    def stream_generate(self, task_key: str, messages, cancel_check=None, request_id=None):
        with closing(get_db()) as db:
            models,routing_set_id=self._models_for_task(db,task_key)
        if not models:
            raise HTTPException(503,"В наборе маршрутизации нет доступных моделей")
        attempts=[]
        with AI_SEMAPHORE:
            for model in models:
                if cancel_check and cancel_check():
                    raise RuntimeError("REQUEST_CANCELLED")
                started=time_monotonic()
                try:
                    api_key=os.getenv(model["api_key_env"]) if model["api_key_env"] else None
                    headers={"Content-Type":"application/json","Accept":"text/event-stream"}
                    if api_key:
                        headers["Authorization"]="Bearer "+api_key
                    payload={"model":model["model_name"],"messages":messages,"stream":True}
                    if model["temperature"] is not None:
                        payload["temperature"]=model["temperature"]
                    if model["max_tokens"] is not None:
                        payload["max_tokens"]=model["max_tokens"]
                    timeout=int(model["timeout"] or REQUEST_TIMEOUT)
                    url=model["base_url"].rstrip("/")+"/chat/completions"
                    req=urllib.request.Request(url,data=json.dumps(payload).encode("utf-8"),headers=headers,method="POST")
                    response=urllib.request.urlopen(req,timeout=timeout)
                    try:
                        saw_content=False
                        for raw_line in response:
                            if cancel_check and cancel_check():
                                raise RuntimeError("REQUEST_CANCELLED")
                            line=raw_line.decode("utf-8",errors="replace").strip()
                            if not line or not line.startswith("data:"):
                                continue
                            data_text=line[5:].strip()
                            if data_text == "[DONE]":
                                break
                            try:
                                data=json.loads(data_text)
                            except json.JSONDecodeError:
                                continue
                            delta=data.get("choices",[{}])[0].get("delta",{}).get("content")
                            if isinstance(delta,str) and delta:
                                saw_content=True
                                yield {"type":"delta","content":delta}
                        if cancel_check and cancel_check():
                            raise RuntimeError("REQUEST_CANCELLED")
                        if not saw_content:
                            raise RuntimeError("Провайдер вернул пустой поток")
                    finally:
                        response.close()
                    attempts.append({"provider":model["provider_name"],"model":model["model_name"],"status":"success","duration_ms":int((time_monotonic()-started)*1000)})
                    yield {"type":"done","model":model["model_name"],"provider":model["provider_name"],"routing_set_id":routing_set_id,"fallback_attempts":attempts}
                    return
                except Exception as exc:
                    attempts.append({"provider":model["provider_name"],"model":model["model_name"],"status":"failed","error":str(exc)[:500],"duration_ms":int((time_monotonic()-started)*1000)})
                    if request_id:
                        update_ai_request(request_id,fallback_attempts_json=json.dumps(attempts,ensure_ascii=False))
                    if request_id and cancel_check and cancel_check():
                        raise RuntimeError("REQUEST_CANCELLED")
                    if model is not models[-1]:
                        yield {"type":"fallback","provider":model["provider_name"],"model":model["model_name"],"error":str(exc)[:200]}
        raise RuntimeError("ALL_MODELS_UNAVAILABLE")

    def generate(self, task_key: str, messages, cancel_check=None, request_id=None):
        with closing(get_db()) as db:
            models,routing_set_id=self._models_for_task(db,task_key)
        if not models:
            raise HTTPException(503,"В наборе маршрутизации нет доступных моделей")
        attempts=[]
        with AI_SEMAPHORE:
            for model in models:
                if cancel_check and cancel_check():
                    raise RuntimeError("REQUEST_CANCELLED")
                started=time_monotonic()
                try:
                    content=self._request(model,messages)
                    if cancel_check and cancel_check():
                        raise RuntimeError("REQUEST_CANCELLED")
                    attempts.append({"provider":model["provider_name"],"model":model["model_name"],"status":"success","duration_ms":int((time_monotonic()-started)*1000)})
                    return {"content":content,"model":model["model_name"],"provider":model["provider_name"],"routing_set_id":routing_set_id,"fallback_attempts":attempts}
                except Exception as exc:
                    attempts.append({"provider":model["provider_name"],"model":model["model_name"],"status":"failed","error":str(exc)[:500],"duration_ms":int((time_monotonic()-started)*1000)})
                    if request_id:
                        update_ai_request(request_id,fallback_attempts_json=json.dumps(attempts,ensure_ascii=False))
                    if request_id and cancel_check and cancel_check():
                        raise RuntimeError("REQUEST_CANCELLED")
        raise RuntimeError("ALL_MODELS_UNAVAILABLE")

def time_monotonic():
    return time.monotonic()

def create_ai_request(user_id:int,chat_id:int,task_key:str,message_id=None):
    ts=now_iso()
    with closing(get_db()) as db:
        try:
            cur=db.execute("INSERT INTO ai_requests(user_id,chat_id,message_id,task_key,status,created_at) VALUES(?,?,?,?,?,?)",(user_id,chat_id,message_id,task_key,"queued",ts))
            db.commit()
        except sqlite3.IntegrityError as exc:
            db.rollback()
            if task_key == "main_generation":
                raise HTTPException(409,"В этом чате уже выполняется запрос")
            raise
        return int(cur.lastrowid)

def is_ai_request_cancelled(request_id:int) -> bool:
    with closing(get_db()) as db:
        row=db.execute("SELECT status,cancel_requested FROM ai_requests WHERE id=?",(request_id,)).fetchone()
    if row is None:
        return False
    return bool(row["cancel_requested"]) or row["status"] == "cancelled"

def update_ai_request(request_id:int,**values):
    if not values:
        return
    allowed={"status","error","provider","model","routing_set","fallback_attempts_json","message_id","started_at","completed_at","cancel_requested"}
    values={k:v for k,v in values.items() if k in allowed}
    if not values: return
    parts=[k+"=?" for k in values]
    params=list(values.values())+[request_id]
    with closing(get_db()) as db:
        db.execute("UPDATE ai_requests SET "+",".join(parts)+" WHERE id=?",params)
        db.commit()

def ai_request_owned(db,request_id:int,user_id:int):
    row=db.execute("SELECT * FROM ai_requests WHERE id=? AND user_id=?",(request_id,user_id)).fetchone()
    if row is None: raise HTTPException(404,"Запрос не найден")
    return row

def ensure_no_active_main_request(db,chat_id:int):
    row=db.execute(
        "SELECT id FROM ai_requests WHERE chat_id=? AND task_key='main_generation' AND status IN ('queued','processing') LIMIT 1",
        (chat_id,)
    ).fetchone()
    if row is not None:
        raise HTTPException(409,"В этом чате уже выполняется запрос")

MODEL_ROUTER=ModelRouter()

def _safe_title_context(messages):
    compact=[]
    for row in messages[-6:]:
        content=str(row.get("content","")).strip()
        if len(content)>1200:
            content=content[:1200]+"…"
        compact.append({"role":row.get("role","user"),"content":content})
    return compact

def _parse_suggestions(raw: str):
    text=raw.strip()
    try:
        value=json.loads(text)
        if isinstance(value,dict):
            value=value.get("suggestions",[])
        if isinstance(value,list):
            items=[str(x).strip() for x in value if str(x).strip()]
            if items:
                return items[:5]
    except Exception:
        pass
    items=[]
    for line in text.splitlines():
        line=line.strip().lstrip("-•*0123456789.) ").strip()
        if line:
            items.append(line)
    return items[:5]

def _run_post_response_tasks(chat_id:int, user_id:int):
    try:
        with closing(get_db()) as db:
            chat=get_owned_chat(db,chat_id,user_id)
            rows=db.execute("SELECT role,content FROM messages WHERE chat_id=? ORDER BY created_at ASC,id ASC",(chat_id,)).fetchall()
            messages=[{"role":row["role"],"content":row["content"]} for row in rows]

        if chat["title"].strip() == "Новый чат":
            title_messages=[
                {"role":"system","content":"Придумай короткое точное название для чата на русском языке. Верни только название, без кавычек и пояснений. Максимум 80 символов."},
                {"role":"user","content":json.dumps(_safe_title_context(messages),ensure_ascii=False)}
            ]
            title_result=MODEL_ROUTER.generate("title_generation",title_messages)
            title=title_result["content"].strip().replace("\n"," ")
            title=title.strip('"«»')[:80].strip()
            if title:
                with closing(get_db()) as db:
                    db.execute("UPDATE chats SET title=?,updated_at=? WHERE id=? AND user_id=? AND title='Новый чат'",(title,now_iso(),chat_id,user_id))
                    db.commit()

        suggestion_messages=[
            {"role":"system","content":"Предложи 3-5 коротких полезных продолжений текущего диалога на русском языке. Верни строго JSON-массив строк без markdown и пояснений."},
            {"role":"user","content":json.dumps(_safe_title_context(messages),ensure_ascii=False)}
        ]
        suggestion_result=MODEL_ROUTER.generate("suggestions_generation",suggestion_messages)
        suggestions=_parse_suggestions(suggestion_result["content"])
        if suggestions:
            with closing(get_db()) as db:
                ts=now_iso()
                db.execute("DELETE FROM chat_suggestions WHERE chat_id=?",(chat_id,))
                db.execute("INSERT INTO chat_suggestions(chat_id,suggestions_json,created_at,updated_at) VALUES(?,?,?,?)",(chat_id,json.dumps(suggestions,ensure_ascii=False),ts,ts))
                db.commit()
    except Exception:
        import logging
        logging.exception("Post-response AI tasks failed for chat_id=%s",chat_id)

app = FastAPI(title="ChatStudio API",version="0.1.0")

@app.middleware("http")
async def same_origin_api_middleware(request: Request, call_next):
    if CSRF_PROTECTION and request.method in {"POST","PUT","PATCH","DELETE"} and request.url.path.startswith("/api/"):
        origin = request.headers.get("origin")
        if origin:
            expected = f"{request.url.scheme}://{request.headers.get('host', '')}"
            if origin.rstrip("/") != expected.rstrip("/"):
                return JSONResponse(status_code=403, content={"detail":"Недопустимый источник запроса"})
    return await call_next(request)
app.add_middleware(SessionMiddleware,secret_key=SESSION_SECRET,session_cookie="chatstudio_session",same_site="lax",https_only=SESSION_COOKIE_SECURE,max_age=60*60*24*30)
app.add_middleware(CORSMiddleware,allow_origins=["*"],allow_credentials=False,allow_methods=["*"],allow_headers=["*"])
init_db()

# A queued request is never actively processed by another task: generation starts
# synchronously in the same HTTP request. If the process restarts after creating
# the record but before generation begins, do not leave that record stuck forever.
with closing(get_db()) as _startup_db:
    _startup_db.execute(
        "UPDATE ai_requests SET status='failed',error=?,completed_at=? "
        "WHERE status IN ('queued','processing')",
        ("Сервис был перезапущен до завершения генерации",now_iso())
    )
    _startup_db.commit()

@app.get("/api/health")
def health():
    return {"status":"ok"}

@app.post("/api/auth/register")
def register(payload:RegisterRequest,request:Request):
    _check_same_origin(request)
    _check_auth_rate(request, "register", AUTH_REGISTER_LIMIT)
    name,email=payload.name.strip(),payload.email.strip().lower()
    if not name or "@" not in email:
        raise HTTPException(400,"Некорректные данные")
    with closing(get_db()) as db:
        if db.execute("SELECT id FROM users WHERE email=?",(email,)).fetchone():
            raise HTTPException(409,"Пользователь уже существует")
        ts=now_iso()
        cur=db.execute("INSERT INTO users(email,password_hash,name,created_at,updated_at) VALUES(?,?,?,?,?)",(email,hash_password(payload.password),name,ts,ts))
        db.commit()
        user_id=int(cur.lastrowid)
    request.session["user_id"]=user_id
    return {"user":{"id":user_id,"email":email,"name":name}}

@app.post("/api/auth/login")
def login(payload:LoginRequest,request:Request):
    _check_same_origin(request)
    _check_auth_rate(request, "login", AUTH_LOGIN_LIMIT)
    email=payload.email.strip().lower()
    with closing(get_db()) as db:
        user=db.execute("SELECT id,email,name,password_hash FROM users WHERE email=?",(email,)).fetchone()
    if user is None or not verify_password(payload.password,user["password_hash"]):
        raise HTTPException(401,"Неверный email или пароль")
    request.session["user_id"]=int(user["id"])
    _clear_auth_rate(request, "login")
    return {"user":{"id":int(user["id"]),"email":user["email"],"name":user["name"]}}

@app.patch("/api/auth/profile")
def update_profile(payload:ProfileUpdateRequest,request:Request):
    _check_same_origin(request)
    user=current_user(request)
    name=payload.name.strip()
    if not name: raise HTTPException(400,"Имя не может быть пустым")
    with closing(get_db()) as db:
        db.execute("UPDATE users SET name=?,updated_at=? WHERE id=?",(name,now_iso(),user["id"]))
        db.commit()
        row=db.execute("SELECT id,email,name,created_at,updated_at FROM users WHERE id=?",(user["id"],)).fetchone()
    return {"user":dict(row)}

@app.patch("/api/auth/password")
def update_password(payload:PasswordUpdateRequest,request:Request):
    user=current_user(request)
    if not verify_password(payload.current_password,user["password_hash"]): raise HTTPException(400,"Текущий пароль указан неверно")
    with closing(get_db()) as db:
        db.execute("UPDATE users SET password_hash=?,updated_at=? WHERE id=?",(hash_password(payload.new_password),now_iso(),user["id"]))
        db.commit()
    return {"ok":True}

@app.post("/api/auth/logout")
def logout(request:Request):
    _check_same_origin(request)
    request.session.clear()
    return {"ok":True}

@app.get("/api/auth/me")
def me(request:Request):
    return {"user":dict(current_user(request))}

@app.get("/api/chats")
def list_chats(request:Request):
    user=current_user(request)
    with closing(get_db()) as db:
        rows=db.execute("SELECT id,title,archived,created_at,updated_at FROM chats WHERE user_id=? ORDER BY updated_at DESC",(user["id"],)).fetchall()
    return {"chats":[dict(x) for x in rows]}

@app.post("/api/chats")
def create_chat(payload:CreateChatRequest,request:Request):
    user=current_user(request)
    title=payload.title.strip() or "Новый чат"
    ts=now_iso()
    with closing(get_db()) as db:
        cur=db.execute("INSERT INTO chats(user_id,title,created_at,updated_at) VALUES(?,?,?,?)",(user["id"],title,ts,ts))
        db.commit()
        chat=get_owned_chat(db,int(cur.lastrowid),int(user["id"]))
    return {"chat":dict(chat)}

@app.get("/api/chats/{chat_id}")
def get_chat(chat_id:int,request:Request):
    user=current_user(request)
    with closing(get_db()) as db:
        chat=get_owned_chat(db,chat_id,int(user["id"]))
        messages=db.execute(
            "SELECT id,chat_id,user_id,role,content,model,provider,routing_set,parent_message_id,created_at FROM messages WHERE chat_id=? ORDER BY created_at ASC,id ASC",
            (chat_id,)
        ).fetchall()
        result=[]
        for message in messages:
            item=dict(message)
            files=db.execute("SELECT f.id,f.filename,f.mime_type,f.size,f.created_at FROM message_files mf JOIN files f ON f.id=mf.file_id WHERE mf.message_id=? AND f.user_id=? ORDER BY f.id",(message["id"],user["id"])).fetchall()
            item["files"]=[dict(x) for x in files]
            result.append(item)
    return {"chat":dict(chat),"messages":result}

@app.post("/api/chats/{chat_id}/share")
def create_chat_share(chat_id:int,request:Request):
    user=current_user(request)
    with closing(get_db()) as db:
        get_owned_chat(db,chat_id,int(user["id"]))
        row=db.execute("SELECT id,chat_id,token,enabled,created_at,updated_at FROM chat_shares WHERE chat_id=?",(chat_id,)).fetchone()
        if row is None:
            ts=now_iso()
            token=secrets.token_urlsafe(32)
            db.execute("INSERT INTO chat_shares(chat_id,token,enabled,created_at,updated_at) VALUES(?,?,?,?,?)",(chat_id,token,1,ts,ts))
            db.commit()
            row=db.execute("SELECT id,chat_id,token,enabled,created_at,updated_at FROM chat_shares WHERE chat_id=?",(chat_id,)).fetchone()
        elif not row["enabled"]:
            db.execute("UPDATE chat_shares SET enabled=1,updated_at=? WHERE chat_id=?",(now_iso(),chat_id))
            db.commit()
            row=db.execute("SELECT id,chat_id,token,enabled,created_at,updated_at FROM chat_shares WHERE chat_id=?",(chat_id,)).fetchone()
    return {"share":{**dict(row),"url":f"/share/{row['token']}"}}

@app.get("/api/chats/{chat_id}/share")
def get_chat_share(chat_id:int,request:Request):
    user=current_user(request)
    with closing(get_db()) as db:
        get_owned_chat(db,chat_id,int(user["id"]))
        row=db.execute("SELECT id,chat_id,token,enabled,created_at,updated_at FROM chat_shares WHERE chat_id=?",(chat_id,)).fetchone()
    if row is None:
        return {"share":None}
    return {"share":{**dict(row),"url":f"/share/{row['token']}"}}

@app.patch("/api/chats/{chat_id}/share")
def update_chat_share(chat_id:int,payload:ShareUpdateRequest,request:Request):
    user=current_user(request)
    with closing(get_db()) as db:
        get_owned_chat(db,chat_id,int(user["id"]))
        row=db.execute("SELECT id,chat_id,token,enabled,created_at,updated_at FROM chat_shares WHERE chat_id=?",(chat_id,)).fetchone()
        if row is None: raise HTTPException(404,"Ссылка ещё не создана")
        db.execute("UPDATE chat_shares SET enabled=?,updated_at=? WHERE chat_id=?",(int(payload.enabled),now_iso(),chat_id))
        db.commit()
        row=db.execute("SELECT id,chat_id,token,enabled,created_at,updated_at FROM chat_shares WHERE chat_id=?",(chat_id,)).fetchone()
    return {"share":{**dict(row),"url":f"/share/{row['token']}"}}

def _shared_chat(token:str):
    with closing(get_db()) as db:
        share=db.execute("SELECT id,chat_id,token,enabled FROM chat_shares WHERE token=?",(token,)).fetchone()
        if share is None or not share["enabled"]: raise HTTPException(404,"Ссылка недействительна или отключена")
        chat=db.execute("SELECT id,title,created_at,updated_at FROM chats WHERE id=?",(share["chat_id"],)).fetchone()
        if chat is None: raise HTTPException(404,"Чат не найден")
        messages=db.execute("SELECT id,role,content,model,provider,routing_set,parent_message_id,created_at FROM messages WHERE chat_id=? ORDER BY created_at ASC,id ASC",(share["chat_id"],)).fetchall()
        result=[]
        for message in messages:
            item=dict(message)
            files=db.execute("SELECT f.id,f.filename,f.mime_type,f.size FROM message_files mf JOIN files f ON f.id=mf.file_id JOIN chats c ON c.id=? WHERE mf.message_id=? AND f.user_id=c.user_id ORDER BY f.id",(share["chat_id"],message["id"])).fetchall()
            item["files"]=[dict(x) for x in files]
            result.append(item)
    return dict(chat),result,share

@app.get("/api/shared/{token}")
def get_shared_chat(token:str):
    chat,messages,share=_shared_chat(token)
    return {"chat":chat,"messages":messages,"read_only":True}

@app.get("/api/shared/{token}/files/{file_id}")
def get_shared_file(token:str,file_id:int):
    chat,_,share=_shared_chat(token)
    with closing(get_db()) as db:
        row=db.execute("SELECT f.path,f.filename,f.mime_type FROM files f JOIN message_files mf ON mf.file_id=f.id JOIN messages m ON m.id=mf.message_id JOIN chats c ON c.id=m.chat_id WHERE f.id=? AND m.chat_id=? AND f.user_id=c.user_id",(file_id,chat["id"])).fetchone()
    if row is None: raise HTTPException(404,"Файл не найден")
    path=_stored_file_path(row["path"])
    if not path.is_file(): raise HTTPException(404,"Файл не найден")
    return FileResponse(str(path),media_type=row["mime_type"] or "application/octet-stream",filename=row["filename"])

@app.patch("/api/chats/{chat_id}")
def rename_chat(chat_id:int,payload:RenameChatRequest,request:Request):
    user=current_user(request)
    title=payload.title.strip()
    if not title:
        raise HTTPException(400,"Название не может быть пустым")
    with closing(get_db()) as db:
        get_owned_chat(db,chat_id,int(user["id"]))
        db.execute("UPDATE chats SET title=?,updated_at=? WHERE id=? AND user_id=?",(title,now_iso(),chat_id,user["id"]))
        db.commit()
        chat=get_owned_chat(db,chat_id,int(user["id"]))
    return {"chat":dict(chat)}

@app.patch("/api/chats/{chat_id}/archive")
def archive_chat(chat_id:int,request:Request):
    user=current_user(request)
    with closing(get_db()) as db:
        chat=get_owned_chat(db,chat_id,int(user["id"]))
        archived=0 if int(chat["archived"]) else 1
        db.execute("UPDATE chats SET archived=?,updated_at=? WHERE id=? AND user_id=?",(archived,now_iso(),chat_id,user["id"]))
        db.commit()
        chat=get_owned_chat(db,chat_id,int(user["id"]))
    return {"chat":dict(chat)}

@app.delete("/api/chats/{chat_id}")
def delete_chat(chat_id:int,request:Request):
    user=current_user(request)
    with closing(get_db()) as db:
        get_owned_chat(db,chat_id,int(user["id"]))
        db.execute("DELETE FROM chats WHERE id=? AND user_id=?",(chat_id,user["id"]))
        db.commit()
    return {"ok":True}

@app.get("/api/search")
def search(request:Request,q:str=""):
    user=current_user(request)
    query=q.strip()[:MAX_SEARCH_LENGTH]
    if not query:
        return {"results":[]}
    like="%" + query + "%"
    with closing(get_db()) as db:
        rows=db.execute(
            "SELECT DISTINCT c.id AS chat_id,c.title AS chat_title,m.id AS message_id,m.role,m.content,m.created_at FROM messages m JOIN chats c ON c.id=m.chat_id LEFT JOIN message_files mf ON mf.message_id=m.id LEFT JOIN files f ON f.id=mf.file_id WHERE c.user_id=? AND (c.title LIKE ? OR m.content LIKE ? OR f.filename LIKE ?) ORDER BY m.created_at DESC LIMIT 50",
            (user["id"],like,like,like)
        ).fetchall()
    return {"results":[dict(x) for x in rows]}

@app.get("/api/chats/{chat_id}/suggestions")
def get_chat_suggestions(chat_id:int,request:Request):
    user=current_user(request)
    with closing(get_db()) as db:
        get_owned_chat(db,chat_id,int(user["id"]))
        row=db.execute("SELECT suggestions_json,updated_at FROM chat_suggestions WHERE chat_id=? ORDER BY updated_at DESC LIMIT 1",(chat_id,)).fetchone()
    if row is None:
        return {"suggestions":[],"updated_at":None}
    try:
        suggestions=json.loads(row["suggestions_json"])
    except Exception:
        suggestions=[]
    return {"suggestions":suggestions,"updated_at":row["updated_at"]}

@app.post("/api/requests")
def create_request(payload:MessageRequest,request:Request,background_tasks:BackgroundTasks):
    user=current_user(request)
    content=payload.content.strip()
    if not content and not payload.file_ids:
        raise HTTPException(400,"Сообщение не может быть пустым")
    with closing(get_db()) as db:
        get_owned_chat(db,payload.chat_id,int(user["id"]))
        if payload.parent_message_id is not None:
            parent=db.execute("SELECT id FROM messages WHERE id=? AND chat_id=? AND user_id=?",(payload.parent_message_id,payload.chat_id,user["id"])).fetchone()
            if parent is None:
                raise HTTPException(404,"Родительское сообщение не найдено")
        if len(set(payload.file_ids)) != len(payload.file_ids):
            raise HTTPException(400,"Файлы указаны повторно")
        files=[]
        if payload.file_ids:
            placeholders=",".join("?" for _ in payload.file_ids)
            files=db.execute(f"SELECT id,filename,mime_type,size,path FROM files WHERE user_id=? AND id IN ({placeholders})",(user["id"],*payload.file_ids)).fetchall()
            if len(files)!=len(payload.file_ids): raise HTTPException(404,"Один или несколько файлов не найдены")
            if sum(int(x["size"]) for x in files)>MAX_TOTAL_FILE_SIZE:
                raise HTTPException(413,f"Общий размер файлов слишком большой. Максимум: {MAX_TOTAL_FILE_SIZE // 1024 // 1024} МБ")
        request_id=create_ai_request(int(user["id"]),payload.chat_id,"main_generation",None)
        ts=now_iso()
        try:
            if payload.edit_message_id is not None:
                target=db.execute("SELECT id FROM messages WHERE id=? AND chat_id=? AND user_id=? AND role='user'",(payload.edit_message_id,payload.chat_id,user["id"])).fetchone()
                if target is None:
                    raise HTTPException(404,"Редактируемое сообщение не найдено")
                if payload.parent_message_id is not None and int(payload.parent_message_id)==int(payload.edit_message_id):
                    raise HTTPException(400,"Сообщение не может быть родителем самого себя")
                db.execute("DELETE FROM messages WHERE chat_id=? AND id>?",(payload.chat_id,payload.edit_message_id))
                db.execute("UPDATE messages SET content=?,parent_message_id=? WHERE id=?",(content,payload.parent_message_id,payload.edit_message_id))
                db.execute("DELETE FROM message_files WHERE message_id=?",(payload.edit_message_id,))
                if files:
                    db.executemany("INSERT INTO message_files(message_id,file_id) VALUES(?,?)",[(payload.edit_message_id,int(x["id"])) for x in files])
                user_message_id=int(payload.edit_message_id)
            else:
                cur=db.execute("INSERT INTO messages(chat_id,user_id,role,content,parent_message_id,created_at) VALUES(?,?, 'user',?,?,?)",(payload.chat_id,user["id"],content,payload.parent_message_id,ts))
                if files:
                    db.executemany("INSERT INTO message_files(message_id,file_id) VALUES(?,?)",[(int(cur.lastrowid),int(x["id"])) for x in files])
                user_message_id=int(cur.lastrowid)
            db.execute("UPDATE chats SET updated_at=? WHERE id=? AND user_id=?",(ts,payload.chat_id,user["id"]))
            db.commit()
        except Exception:
            db.rollback()
            update_ai_request(request_id,status="failed",error="Не удалось сохранить сообщение",completed_at=now_iso())
            raise
    update_ai_request(request_id,message_id=user_message_id,status="processing",started_at=now_iso())
    with closing(get_db()) as db:
        ai_messages=_chat_ai_messages(db,payload.chat_id)
    try:
        result=MODEL_ROUTER.generate("main_generation",ai_messages,cancel_check=lambda: is_ai_request_cancelled(request_id),request_id=request_id)
    except HTTPException as exc:
        update_ai_request(request_id,status="failed",error=str(exc.detail),completed_at=now_iso())
        raise
    except RuntimeError as exc:
        status="cancelled" if str(exc)=="REQUEST_CANCELLED" else "failed"
        update_ai_request(request_id,status=status,error=str(exc),completed_at=now_iso())
        raise HTTPException(502,"Запрос не выполнен")
    except Exception as exc:
        update_ai_request(request_id,status="failed",error=str(exc)[:500],completed_at=now_iso())
        raise HTTPException(502,"Запрос не выполнен")
    try:
        with closing(get_db()) as db:
            ts=now_iso()
            cur=db.execute("INSERT INTO messages(chat_id,user_id,role,content,model,provider,routing_set,parent_message_id,created_at) VALUES(?,?, 'assistant',?,?,?,?,?,?,?)",(payload.chat_id,user["id"],result["content"],result["model"],result["provider"],str(result["routing_set_id"]),user_message_id,ts))
            db.execute("UPDATE chats SET updated_at=? WHERE id=? AND user_id=?",(ts,payload.chat_id,user["id"]))
            completion=db.execute("UPDATE ai_requests SET status='completed',message_id=?,provider=?,model=?,routing_set=?,fallback_attempts_json=?,completed_at=? WHERE id=? AND status='processing' AND cancel_requested=0",(int(cur.lastrowid),result["provider"],result["model"],str(result["routing_set_id"]),json.dumps(result.get("fallback_attempts",[]),ensure_ascii=False),ts,request_id))
            if completion.rowcount != 1:
                db.rollback()
                raise RuntimeError("REQUEST_CANCELLED")
            db.commit()
            assistant=db.execute("SELECT id,chat_id,user_id,role,content,model,provider,routing_set,parent_message_id,created_at FROM messages WHERE id=?",(cur.lastrowid,)).fetchone()
            user_message=db.execute("SELECT id,chat_id,user_id,role,content,model,provider,routing_set,parent_message_id,created_at FROM messages WHERE id=?",(user_message_id,)).fetchone()
    except RuntimeError as exc:
        status="cancelled" if str(exc)=="REQUEST_CANCELLED" else "failed"
        update_ai_request(request_id,status=status,error=str(exc),completed_at=now_iso())
        raise HTTPException(409 if status=="cancelled" else 502,"Запрос остановлен" if status=="cancelled" else "Не удалось сохранить ответ")
    except Exception as exc:
        update_ai_request(request_id,status="failed",error=str(exc)[:500],completed_at=now_iso())
        raise HTTPException(502,"Не удалось сохранить ответ")
    background_tasks.add_task(_run_post_response_tasks,payload.chat_id,int(user["id"]))
    return {"status":"completed","request_id":request_id,"message":dict(user_message),"assistant":dict(assistant)}
 
class RoutingSetUpdateRequest(BaseModel):
    name: str | None = Field(default=None,min_length=1,max_length=120)
    description: str | None = Field(default=None,max_length=500)

class RoutingSetPriorityRequest(BaseModel):
    priority: int = Field(ge=1)

def _normalize_routing_priorities(db, routing_set_id):
    rows=db.execute("SELECT model_config_id FROM routing_set_models WHERE routing_set_id=? ORDER BY priority,model_config_id",(routing_set_id,)).fetchall()
    for priority,row in enumerate(rows,1):
        db.execute("UPDATE routing_set_models SET priority=? WHERE routing_set_id=? AND model_config_id=?",(priority+1000,routing_set_id,row["model_config_id"]))
    for priority,row in enumerate(rows,1):
        db.execute("UPDATE routing_set_models SET priority=? WHERE routing_set_id=? AND model_config_id=?",(priority,routing_set_id,row["model_config_id"]))

@app.patch("/api/routing/sets/{routing_set_id}")
def update_routing_set(routing_set_id:int,payload:RoutingSetUpdateRequest,request:Request):
    require_admin(request)
    with closing(get_db()) as db:
        row=db.execute("SELECT id,name,description,created_at,updated_at FROM routing_sets WHERE id=?",(routing_set_id,)).fetchone()
        if row is None: raise HTTPException(404,"Набор маршрутизации не найден")
        name=payload.name.strip() if payload.name is not None else row["name"]
        description=payload.description if payload.description is not None else row["description"]
        try:
            db.execute("UPDATE routing_sets SET name=?,description=?,updated_at=? WHERE id=?",(name,description,now_iso(),routing_set_id))
            db.commit()
        except sqlite3.IntegrityError: raise HTTPException(409,"Набор с таким названием уже существует")
        updated=db.execute("SELECT id,name,description,created_at,updated_at FROM routing_sets WHERE id=?",(routing_set_id,)).fetchone()
    return {"routing_set":dict(updated)}

@app.delete("/api/routing/sets/{routing_set_id}")
def delete_routing_set(routing_set_id:int,request:Request):
    require_admin(request)
    with closing(get_db()) as db:
        if db.execute("SELECT id FROM routing_sets WHERE id=?",(routing_set_id,)).fetchone() is None:
            raise HTTPException(404,"Набор маршрутизации не найден")
        if db.execute("SELECT task_key FROM task_routes WHERE routing_set_id=?",(routing_set_id,)).fetchone():
            raise HTTPException(409,"Нельзя удалить набор, пока он назначен AI-задаче")
        db.execute("DELETE FROM routing_sets WHERE id=?",(routing_set_id,))
        db.commit()
    return {"ok":True}

@app.delete("/api/routing/sets/{routing_set_id}/models/{model_config_id}")
def remove_model_from_routing_set(routing_set_id:int,model_config_id:int,request:Request):
    require_admin(request)
    with closing(get_db()) as db:
        if db.execute("SELECT id FROM routing_sets WHERE id=?",(routing_set_id,)).fetchone() is None:
            raise HTTPException(404,"Набор маршрутизации не найден")
        if not db.execute("SELECT 1 FROM routing_set_models WHERE routing_set_id=? AND model_config_id=?",(routing_set_id,model_config_id)).fetchone():
            raise HTTPException(404,"Модель не находится в этом наборе")
        db.execute("DELETE FROM routing_set_models WHERE routing_set_id=? AND model_config_id=?",(routing_set_id,model_config_id))
        _normalize_routing_priorities(db,routing_set_id)
        db.execute("UPDATE routing_sets SET updated_at=? WHERE id=?",(now_iso(),routing_set_id))
        db.commit()
    return {"ok":True}

@app.patch("/api/routing/sets/{routing_set_id}/models/{model_config_id}")
def change_routing_model_priority(routing_set_id:int,model_config_id:int,payload:RoutingSetPriorityRequest,request:Request):
    require_admin(request)
    with closing(get_db()) as db:
        rows=db.execute("SELECT model_config_id FROM routing_set_models WHERE routing_set_id=? ORDER BY priority,model_config_id",(routing_set_id,)).fetchall()
        ids=[int(x["model_config_id"]) for x in rows]
        if model_config_id not in ids: raise HTTPException(404,"Модель не находится в этом наборе")
        if payload.priority>len(ids): raise HTTPException(400,"Приоритет вне диапазона")
        ids.remove(model_config_id)
        ids.insert(payload.priority-1,model_config_id)
        for priority,mid in enumerate(ids,1):
            db.execute("UPDATE routing_set_models SET priority=? WHERE routing_set_id=? AND model_config_id=?",(priority+1000,routing_set_id,mid))
        for priority,mid in enumerate(ids,1):
            db.execute("UPDATE routing_set_models SET priority=? WHERE routing_set_id=? AND model_config_id=?",(priority,routing_set_id,mid))
        db.execute("UPDATE routing_sets SET updated_at=? WHERE id=?",(now_iso(),routing_set_id))
        db.commit()
    return {"ok":True}

@app.post("/api/requests/{request_id}/cancel")
def cancel_request(request_id:int,request:Request):
    user=current_user(request)
    with closing(get_db()) as db:
        row=ai_request_owned(db,request_id,int(user["id"]))
        status=row["status"]
        if status in {"completed","failed","cancelled"}:
            data=dict(row)
        elif status == "queued":
            ts=now_iso()
            db.execute(
                "UPDATE ai_requests SET status='cancelled',cancel_requested=1,completed_at=? "
                "WHERE id=? AND user_id=? AND status='queued' AND cancel_requested=0",
                (ts,request_id,user["id"])
            )
            db.commit()
            data=dict(ai_request_owned(db,request_id,int(user["id"])))
        else:
            db.execute(
                "UPDATE ai_requests SET cancel_requested=1 WHERE id=? AND user_id=?",
                (request_id,user["id"])
            )
            db.commit()
            data=dict(ai_request_owned(db,request_id,int(user["id"])))
    return {
        "request_id":request_id,
        "status":data["status"],
        "cancel_requested":bool(data["cancel_requested"]),
        "request":data
    }


@app.post("/api/messages/{message_id}/edit")
def edit_message(message_id:int,payload:MessageRequest,request:Request):
    user=current_user(request)
    with closing(get_db()) as db:
        row=db.execute("SELECT m.*,c.user_id FROM messages m JOIN chats c ON c.id=m.chat_id WHERE m.id=? AND m.role='user'",(message_id,)).fetchone()
        if row is not None:
            ensure_no_active_main_request(db,int(row["chat_id"]))
        if row is None or int(row["user_id"])!=int(user["id"]): raise HTTPException(404,"Сообщение не найдено")
        if not payload.content and not payload.file_ids: raise HTTPException(400,"Сообщение не может быть пустым")
        if payload.parent_message_id is not None:
            parent=db.execute("SELECT id FROM messages WHERE id=? AND chat_id=? AND user_id=?",(payload.parent_message_id,row["chat_id"],user["id"])).fetchone()
            if parent is None: raise HTTPException(404,"Родительское сообщение не найдено")
        files=[]
        if payload.file_ids:
            placeholders=",".join("?" for _ in payload.file_ids)
            files=db.execute(f"SELECT id FROM files WHERE user_id=? AND id IN ({placeholders})",(user["id"],*payload.file_ids)).fetchall()
            if len(files)!=len(set(payload.file_ids)): raise HTTPException(404,"Файл не найден")
        db.execute("UPDATE messages SET content=?,parent_message_id=? WHERE id=?",(payload.content,payload.parent_message_id,message_id))
        db.execute("DELETE FROM message_files WHERE message_id=?",(message_id,))
        if files:
            db.executemany("INSERT INTO message_files(message_id,file_id) VALUES(?,?)",[(message_id,int(x["id"])) for x in files])
        db.execute("UPDATE chats SET updated_at=? WHERE id=?",(now_iso(),row["chat_id"]))
        db.commit()
    return {"ok":True,"chat_id":row["chat_id"],"message_id":message_id}

@app.post("/api/messages/{message_id}/branch")
def branch_message(message_id:int,request:Request):
    user=current_user(request)
    with closing(get_db()) as db:
        row=db.execute("SELECT m.*,c.user_id,c.title FROM messages m JOIN chats c ON c.id=m.chat_id WHERE m.id=?",(message_id,)).fetchone()
        if row is None or int(row["user_id"])!=int(user["id"]):
            raise HTTPException(404,"Сообщение не найдено")
        ensure_no_active_main_request(db,int(row["chat_id"]))
        ts=now_iso()
        title=(row["title"] or "Новый чат")+" — ветка"
        cur=db.execute("INSERT INTO chats(user_id,title,created_at,updated_at,archived) VALUES(?,?,?,?,0)",(user["id"],title,ts,ts))
        new_chat_id=int(cur.lastrowid)
        rows=db.execute("SELECT * FROM messages WHERE chat_id=? AND id<=? ORDER BY created_at ASC,id ASC",(row["chat_id"],message_id)).fetchall()
        idmap={}
        for old in rows:
            new_parent=idmap.get(old["parent_message_id"])
            c=db.execute("INSERT INTO messages(chat_id,user_id,role,content,model,provider,routing_set,parent_message_id,created_at) VALUES(?,?,?,?,?,?,?,?,?)",(new_chat_id,user["id"],old["role"],old["content"],old["model"],old["provider"],old["routing_set"],new_parent,old["created_at"]))
            idmap[int(old["id"])]=int(c.lastrowid)
            files=db.execute(
                "SELECT mf.file_id FROM message_files mf JOIN files f ON f.id=mf.file_id "
                "WHERE mf.message_id=? AND f.user_id=?",
                (old["id"],user["id"])
            ).fetchall()
            db.executemany(
                "INSERT INTO message_files(message_id,file_id) VALUES(?,?)",
                [(int(c.lastrowid),int(x["file_id"])) for x in files]
            )
        db.commit()
    return {"chat_id":new_chat_id}

@app.post("/api/messages/{message_id}/retry")
def retry_message(message_id:int,request:Request):
    user=current_user(request)
    with closing(get_db()) as db:
        row=db.execute("SELECT m.*,c.user_id FROM messages m JOIN chats c ON c.id=m.chat_id WHERE m.id=? AND m.role='assistant'",(message_id,)).fetchone()
        if row is None or int(row["user_id"])!=int(user["id"]): raise HTTPException(404,"Ответ не найден")
        parent=db.execute("SELECT * FROM messages WHERE id=? AND chat_id=? AND user_id=? AND role='user'",(row["parent_message_id"],row["chat_id"],user["id"])).fetchone() if row["parent_message_id"] is not None else db.execute("SELECT * FROM messages WHERE chat_id=? AND id<? AND user_id=? AND role='user' ORDER BY id DESC LIMIT 1",(row["chat_id"],row["id"],user["id"])).fetchone()
        if parent is None or parent["role"]!="user": raise HTTPException(400,"Перед ответом не найден запрос пользователя")
        files=db.execute("SELECT file_id FROM message_files WHERE message_id=? ORDER BY file_id",(parent["id"],)).fetchall()
        payload={"chat_id":int(row["chat_id"]),"content":parent["content"],"parent_message_id":int(parent["id"]),"file_ids":[int(x["file_id"]) for x in files]}
    return {"chat_id":payload["chat_id"],"content":payload["content"],"parent_message_id":payload["parent_message_id"],"file_ids":payload["file_ids"]}

@app.post("/api/requests/stream")
def stream_request(payload:MessageRequest,request:Request,background_tasks:BackgroundTasks):
    user=current_user(request)
    content=payload.content.strip()
    if not content and not payload.file_ids:
        raise HTTPException(400,"Сообщение не может быть пустым")
    with closing(get_db()) as db:
        get_owned_chat(db,payload.chat_id,int(user["id"]))
        if payload.parent_message_id is not None:
            parent=db.execute("SELECT id FROM messages WHERE id=? AND chat_id=? AND user_id=?",(payload.parent_message_id,payload.chat_id,user["id"])).fetchone()
            if parent is None:
                raise HTTPException(404,"Родительское сообщение не найдено")
        if len(set(payload.file_ids)) != len(payload.file_ids):
            raise HTTPException(400,"Файлы указаны повторно")
        files=[]
        if payload.file_ids:
            placeholders=",".join("?" for _ in payload.file_ids)
            files=db.execute(f"SELECT id,filename,mime_type,size,path FROM files WHERE user_id=? AND id IN ({placeholders})",(user["id"],*payload.file_ids)).fetchall()
            if len(files)!=len(payload.file_ids): raise HTTPException(404,"Один или несколько файлов не найдены")
            if sum(int(x["size"]) for x in files)>MAX_TOTAL_FILE_SIZE:
                raise HTTPException(413,f"Общий размер файлов слишком большой. Максимум: {MAX_TOTAL_FILE_SIZE // 1024 // 1024} МБ")
        request_id=create_ai_request(int(user["id"]),payload.chat_id,"main_generation",None)
        ts=now_iso()
        try:
            if payload.edit_message_id is not None:
                target=db.execute("SELECT id FROM messages WHERE id=? AND chat_id=? AND user_id=? AND role='user'",(payload.edit_message_id,payload.chat_id,user["id"])).fetchone()
                if target is None:
                    raise HTTPException(404,"Редактируемое сообщение не найдено")
                if payload.parent_message_id is not None and int(payload.parent_message_id)==int(payload.edit_message_id):
                    raise HTTPException(400,"Сообщение не может быть родителем самого себя")
                db.execute("DELETE FROM messages WHERE chat_id=? AND id>?",(payload.chat_id,payload.edit_message_id))
                db.execute("UPDATE messages SET content=?,parent_message_id=? WHERE id=?",(content,payload.parent_message_id,payload.edit_message_id))
                db.execute("DELETE FROM message_files WHERE message_id=?",(payload.edit_message_id,))
                if files:
                    db.executemany("INSERT INTO message_files(message_id,file_id) VALUES(?,?)",[(payload.edit_message_id,int(x["id"])) for x in files])
                user_message_id=int(payload.edit_message_id)
            else:
                cur=db.execute("INSERT INTO messages(chat_id,user_id,role,content,parent_message_id,created_at) VALUES(?,?, 'user',?,?,?)",(payload.chat_id,user["id"],content,payload.parent_message_id,ts))
                if files:
                    db.executemany("INSERT INTO message_files(message_id,file_id) VALUES(?,?)",[(int(cur.lastrowid),int(x["id"])) for x in files])
                user_message_id=int(cur.lastrowid)
            db.execute("UPDATE chats SET updated_at=? WHERE id=? AND user_id=?",(ts,payload.chat_id,user["id"]))
            db.commit()
        except Exception:
            db.rollback()
            update_ai_request(request_id,status="failed",error="Не удалось сохранить сообщение",completed_at=now_iso())
            raise
    update_ai_request(request_id,message_id=user_message_id,status="processing",started_at=now_iso())
    with closing(get_db()) as db:
        ai_messages=_chat_ai_messages(db,payload.chat_id)

    def event(payload_data):
        return "data: "+json.dumps(payload_data,ensure_ascii=False,separators=(",",":"))+"\\n\\n"

    def generate_events():
        full_content=[]
        selected=None
        try:
            yield event({"type":"start","request_id":request_id})
            for item in MODEL_ROUTER.stream_generate("main_generation",ai_messages,cancel_check=lambda: is_ai_request_cancelled(request_id),request_id=request_id):
                item_type=item.get("type")
                if item_type=="delta":
                    full_content.append(item["content"])
                    yield event(item)
                elif item_type=="fallback":
                    # A failed streaming model may already have emitted partial text.
                    # The fallback response must replace it, not append to it.
                    full_content.clear()
                    yield event(item)
                elif item_type=="done":
                    selected=item
            final_content="".join(full_content)
            if not selected or not final_content.strip():
                raise RuntimeError("Провайдер вернул пустой ответ")
            with closing(get_db()) as db:
                ts=now_iso()
                cur=db.execute("INSERT INTO messages(chat_id,user_id,role,content,model,provider,routing_set,parent_message_id,created_at) VALUES(?,?, 'assistant',?,?,?,?,?,?,?)",(payload.chat_id,user["id"],final_content,selected["model"],selected["provider"],str(selected["routing_set_id"]),user_message_id,ts))
                db.execute("UPDATE chats SET updated_at=? WHERE id=? AND user_id=?",(ts,payload.chat_id,user["id"]))
                completion=db.execute(
                    "UPDATE ai_requests SET status='completed',message_id=?,provider=?,model=?,routing_set=?,fallback_attempts_json=?,completed_at=? "
                    "WHERE id=? AND status='processing' AND cancel_requested=0",
                    (int(cur.lastrowid),selected["provider"],selected["model"],str(selected["routing_set_id"]),json.dumps(selected.get("fallback_attempts",[]),ensure_ascii=False),ts,request_id)
                )
                if completion.rowcount != 1:
                    db.rollback()
                    raise RuntimeError("REQUEST_CANCELLED")
                db.commit()
                assistant=db.execute("SELECT id,chat_id,user_id,role,content,model,provider,routing_set,parent_message_id,created_at FROM messages WHERE id=?",(cur.lastrowid,)).fetchone()
                user_message=db.execute("SELECT id,chat_id,user_id,role,content,model,provider,routing_set,parent_message_id,created_at FROM messages WHERE id=?",(user_message_id,)).fetchone()
            background_tasks.add_task(_run_post_response_tasks,payload.chat_id,int(user["id"]))
            yield event({"type":"complete","request_id":request_id,"message":dict(assistant),"user_message":dict(user_message)})
        except RuntimeError as exc:
            status="cancelled" if str(exc)=="REQUEST_CANCELLED" else "failed"
            update_ai_request(request_id,status=status,error=str(exc),completed_at=now_iso())
            yield event({"type":"cancelled" if status=="cancelled" else "error","request_id":request_id,"error":str(exc)})
        except Exception as exc:
            update_ai_request(request_id,status="failed",error=str(exc)[:500],completed_at=now_iso())
            yield event({"type":"error","request_id":request_id,"error":"Запрос не выполнен"})

    return StreamingResponse(generate_events(),media_type="text/event-stream",headers={"Cache-Control":"no-cache","X-Accel-Buffering":"no"})

@app.get("/api/requests/{request_id}")
def get_request_status(request_id:int,request:Request):
    user=current_user(request)
    with closing(get_db()) as db:
        row=ai_request_owned(db,request_id,int(user["id"]))
    data=dict(row)
    data["cancel_requested"]=bool(data.get("cancel_requested"))
    if data.get("fallback_attempts_json"):
        try: data["fallback_attempts"]=json.loads(data["fallback_attempts_json"])
        except Exception: data["fallback_attempts"]=[]
    else: data["fallback_attempts"]=[]
    data.pop("fallback_attempts_json",None)
    return {"request":data}

@app.get("/api/chats/{chat_id}/requests")
def list_chat_requests(chat_id:int,request:Request):
    user=current_user(request)
    with closing(get_db()) as db:
        get_owned_chat(db,chat_id,int(user["id"]))
        rows=db.execute("SELECT * FROM ai_requests WHERE chat_id=? AND user_id=? ORDER BY created_at DESC LIMIT 50",(chat_id,user["id"])).fetchall()
    return {"requests":[dict(x) for x in rows]}

@app.get("/api/routing/access")
def routing_access(request:Request):
    user=current_user(request)
    return {"is_admin": bool(ADMIN_EMAILS and str(user["email"]).lower() in ADMIN_EMAILS)}

@app.get("/api/routing/providers")
def list_providers(request:Request):
    require_admin(request)
    with closing(get_db()) as db:
        rows=db.execute("SELECT id,name,base_url,api_key_env,enabled,created_at,updated_at FROM providers ORDER BY name COLLATE NOCASE").fetchall()
    return {"providers":[dict(x) for x in rows]}

@app.post("/api/routing/providers")
def create_provider(payload:ProviderRequest,request:Request):
    require_admin(request)
    name=payload.name.strip()
    base_url=payload.base_url.strip().rstrip("/")
    api_key_env=payload.api_key_env.strip() if payload.api_key_env else None
    if api_key_env and not __import__("re").fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", api_key_env):
        raise HTTPException(400,"Некорректное имя переменной API-ключа")
    if not name or not base_url:
        raise HTTPException(400,"Название и URL провайдера обязательны")
    if not (base_url.startswith("http://") or base_url.startswith("https://")):
        raise HTTPException(400,"Base URL должен начинаться с http:// или https://")
    ts=now_iso()
    with closing(get_db()) as db:
        try:
            cur=db.execute("INSERT INTO providers(name,base_url,api_key_env,created_at,updated_at) VALUES(?,?,?,?,?)",(name,base_url,api_key_env,ts,ts))
            db.commit()
        except sqlite3.IntegrityError:
            raise HTTPException(409,"Провайдер с таким названием уже существует")
        row=db.execute("SELECT id,name,base_url,api_key_env,enabled,created_at,updated_at FROM providers WHERE id=?",(cur.lastrowid,)).fetchone()
    return {"provider":dict(row)}

@app.patch("/api/routing/providers/{provider_id}")
def update_provider(provider_id:int,payload:ProviderUpdateRequest,request:Request):
    require_admin(request)
    with closing(get_db()) as db:
        row=db.execute("SELECT id,name,base_url,api_key_env,enabled,created_at,updated_at FROM providers WHERE id=?",(provider_id,)).fetchone()
        if row is None: raise HTTPException(404,"Провайдер не найден")
        name=payload.name.strip() if payload.name is not None else row["name"]
        base_url=payload.base_url.strip() if payload.base_url is not None else row["base_url"]
        api_key_env=payload.api_key_env if payload.api_key_env is not None else row["api_key_env"]
        enabled=int(payload.enabled) if payload.enabled is not None else row["enabled"]
        if not name or not base_url: raise HTTPException(400,"Название и URL провайдера обязательны")
        base_url=base_url.rstrip("/")
        if not (base_url.startswith("http://") or base_url.startswith("https://")): raise HTTPException(400,"Base URL должен начинаться с http:// или https://")
        if api_key_env is not None:
            api_key_env=api_key_env.strip() or None
            if api_key_env and not __import__("re").fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", api_key_env):
                raise HTTPException(400,"Некорректное имя переменной API-ключа")
        try:
            db.execute("UPDATE providers SET name=?,base_url=?,api_key_env=?,enabled=?,updated_at=? WHERE id=?",(name,base_url,api_key_env,enabled,now_iso(),provider_id))
            db.commit()
        except sqlite3.IntegrityError: raise HTTPException(409,"Провайдер с таким названием уже существует")
        updated=db.execute("SELECT id,name,base_url,api_key_env,enabled,created_at,updated_at FROM providers WHERE id=?",(provider_id,)).fetchone()
    return {"provider":dict(updated)}

@app.delete("/api/routing/providers/{provider_id}")
def delete_provider(provider_id:int,request:Request):
    require_admin(request)
    with closing(get_db()) as db:
        row=db.execute("SELECT id FROM providers WHERE id=?",(provider_id,)).fetchone()
        if row is None: raise HTTPException(404,"Провайдер не найден")
        if db.execute("SELECT 1 FROM model_configs WHERE provider_id=?",(provider_id,)).fetchone():
            raise HTTPException(409,"Нельзя удалить провайдера, пока у него есть модели")
        db.execute("DELETE FROM providers WHERE id=?",(provider_id,))
        db.commit()
    return {"ok":True}

@app.get("/api/routing/models")
def list_models(request:Request):
    require_admin(request)
    with closing(get_db()) as db:
        rows=db.execute("SELECT mc.id,mc.name,mc.model_name,mc.temperature,mc.max_tokens,mc.timeout,mc.enabled,p.id AS provider_id,p.name AS provider_name FROM model_configs mc JOIN providers p ON p.id=mc.provider_id ORDER BY p.name COLLATE NOCASE,mc.name COLLATE NOCASE").fetchall()
    return {"models":[dict(x) for x in rows]}

@app.patch("/api/routing/models/{model_id}")
def update_model(model_id:int,payload:ModelConfigUpdateRequest,request:Request):
    require_admin(request)
    with closing(get_db()) as db:
        row=db.execute("SELECT id,provider_id,name,model_name,temperature,max_tokens,timeout,enabled,created_at,updated_at FROM model_configs WHERE id=?",(model_id,)).fetchone()
        if row is None: raise HTTPException(404,"Модель не найдена")
        name=payload.name.strip() if payload.name is not None else row["name"]
        model_name=payload.model_name.strip() if payload.model_name is not None else row["model_name"]
        temperature=payload.temperature if payload.temperature is not None else row["temperature"]
        max_tokens=payload.max_tokens if payload.max_tokens is not None else row["max_tokens"]
        timeout=payload.timeout if payload.timeout is not None else row["timeout"]
        enabled=int(payload.enabled) if payload.enabled is not None else row["enabled"]
        if not name or not model_name: raise HTTPException(400,"Название и model_name обязательны")
        try:
            db.execute("UPDATE model_configs SET name=?,model_name=?,temperature=?,max_tokens=?,timeout=?,enabled=?,updated_at=? WHERE id=?",(name,model_name,temperature,max_tokens,timeout,enabled,now_iso(),model_id))
            db.commit()
        except sqlite3.IntegrityError: raise HTTPException(409,"Такая модель уже существует у этого провайдера")
        updated=db.execute("SELECT mc.id,mc.name,mc.model_name,mc.temperature,mc.max_tokens,mc.timeout,mc.enabled,p.id AS provider_id,p.name AS provider_name FROM model_configs mc JOIN providers p ON p.id=mc.provider_id WHERE mc.id=?",(model_id,)).fetchone()
    return {"model":dict(updated)}

@app.delete("/api/routing/models/{model_id}")
def delete_model(model_id:int,request:Request):
    require_admin(request)
    with closing(get_db()) as db:
        if db.execute("SELECT 1 FROM model_configs WHERE id=?",(model_id,)).fetchone() is None:
            raise HTTPException(404,"Модель не найдена")
        if db.execute("SELECT 1 FROM routing_set_models WHERE model_config_id=?",(model_id,)).fetchone():
            raise HTTPException(409,"Нельзя удалить модель, пока она используется в routing set")
        db.execute("DELETE FROM model_configs WHERE id=?",(model_id,))
        db.commit()
    return {"ok":True}

@app.post("/api/routing/models")
def create_model(payload:ModelConfigRequest,request:Request):
    require_admin(request)
    ts=now_iso()
    with closing(get_db()) as db:
        if db.execute("SELECT id FROM providers WHERE id=?",(payload.provider_id,)).fetchone() is None:
            raise HTTPException(404,"Провайдер не найден")
        try:
            cur=db.execute("INSERT INTO model_configs(provider_id,name,model_name,temperature,max_tokens,timeout,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?)",(payload.provider_id,payload.name.strip(),payload.model_name.strip(),payload.temperature,payload.max_tokens,payload.timeout,ts,ts))
            db.commit()
        except sqlite3.IntegrityError:
            raise HTTPException(409,"Такая модель уже существует у этого провайдера")
        row=db.execute("SELECT mc.id,mc.name,mc.model_name,mc.temperature,mc.max_tokens,mc.timeout,mc.enabled,p.id AS provider_id,p.name AS provider_name FROM model_configs mc JOIN providers p ON p.id=mc.provider_id WHERE mc.id=?",(cur.lastrowid,)).fetchone()
    return {"model":dict(row)}

@app.post("/api/routing/sets/{routing_set_id}/models")
def add_model_to_routing_set(routing_set_id:int,payload:RoutingSetModelRequest,request:Request):
    require_admin(request)
    with closing(get_db()) as db:
        if db.execute("SELECT id FROM routing_sets WHERE id=?",(routing_set_id,)).fetchone() is None:
            raise HTTPException(404,"Набор маршрутизации не найден")
        if db.execute("SELECT id FROM model_configs WHERE id=? AND enabled=1",(payload.model_config_id,)).fetchone() is None:
            raise HTTPException(404,"Модель не найдена или отключена")
        if db.execute("SELECT 1 FROM routing_set_models WHERE routing_set_id=? AND priority=?",(routing_set_id,payload.priority)).fetchone():
            raise HTTPException(409,"Этот приоритет уже занят")
        try:
            db.execute("INSERT INTO routing_set_models(routing_set_id,model_config_id,priority) VALUES(?,?,?)",(routing_set_id,payload.model_config_id,payload.priority))
            db.commit()
        except sqlite3.IntegrityError:
            raise HTTPException(409,"Модель уже находится в этом наборе")
    return {"ok":True}

@app.delete("/api/routing/sets/{routing_set_id}/models/{model_config_id}")
def remove_model_from_routing_set(routing_set_id:int,model_config_id:int,request:Request):
    require_admin(request)
    with closing(get_db()) as db:
        db.execute("DELETE FROM routing_set_models WHERE routing_set_id=? AND model_config_id=?",(routing_set_id,model_config_id));db.commit()
    return {"ok":True}

@app.put("/api/routing/tasks/{task_key}")
def set_task_route(task_key:str,payload:TaskRouteRequest,request:Request):
    require_admin(request)
    if task_key not in {"main_generation","title_generation","suggestions_generation"}:
        raise HTTPException(400,"Неизвестная AI-задача")
    with closing(get_db()) as db:
        if payload.routing_set_id is not None:
            if db.execute("SELECT id FROM routing_sets WHERE id=?",(payload.routing_set_id,)).fetchone() is None:
                raise HTTPException(404,"Набор маршрутизации не найден")
            if db.execute(
                "SELECT 1 FROM routing_set_models rsm "
                "JOIN model_configs mc ON mc.id=rsm.model_config_id "
                "JOIN providers p ON p.id=mc.provider_id "
                "WHERE rsm.routing_set_id=? AND mc.enabled=1 AND p.enabled=1 LIMIT 1",
                (payload.routing_set_id,)
            ).fetchone() is None:
                raise HTTPException(409,"Нельзя назначить пустой набор: в нём нет доступной модели")
        db.execute("INSERT INTO task_routes(task_key,routing_set_id,updated_at) VALUES(?,?,?) ON CONFLICT(task_key) DO UPDATE SET routing_set_id=excluded.routing_set_id,updated_at=excluded.updated_at",(task_key,payload.routing_set_id,now_iso()))
        db.commit()
    return {"ok":True,"task_key":task_key,"routing_set_id":payload.routing_set_id}

@app.get("/api/routing/sets")
def list_routing_sets(request:Request):
    require_admin(request)
    with closing(get_db()) as db:
        sets=db.execute("SELECT id,name,description,created_at,updated_at FROM routing_sets ORDER BY name COLLATE NOCASE").fetchall()
        result=[]
        for rs in sets:
            models=db.execute(
                "SELECT rsm.priority,mc.id AS model_id,mc.name AS model_name,mc.model_name AS provider_model,p.id AS provider_id,p.name AS provider_name FROM routing_set_models rsm JOIN model_configs mc ON mc.id=rsm.model_config_id JOIN providers p ON p.id=mc.provider_id WHERE rsm.routing_set_id=? ORDER BY rsm.priority",
                (rs["id"],)
            ).fetchall()
            result.append({**dict(rs),"models":[dict(x) for x in models]})
    return {"routing_sets":result}

@app.post("/api/routing/sets")
def create_routing_set(payload:RoutingSetRequest,request:Request):
    require_admin(request)
    name=payload.name.strip()
    with closing(get_db()) as db:
        try:
            ts=now_iso()
            cur=db.execute("INSERT INTO routing_sets(name,description,created_at,updated_at) VALUES(?,?,?,?)",(name,payload.description,ts,ts))
            db.commit()
        except sqlite3.IntegrityError:
            raise HTTPException(409,"Набор с таким названием уже существует")
        row=db.execute("SELECT id,name,description,created_at,updated_at FROM routing_sets WHERE id=?",(cur.lastrowid,)).fetchone()
    return {"routing_set":dict(row)}

@app.get("/api/routing/tasks/{task_key}")
def get_task_route(task_key:str,request:Request):
    require_admin(request)
    if task_key not in {"main_generation","title_generation","suggestions_generation"}:
        raise HTTPException(400,"Неизвестная AI-задача")
    with closing(get_db()) as db:
        row=db.execute("SELECT tr.task_key,tr.routing_set_id,rs.name AS routing_set_name,tr.updated_at FROM task_routes tr LEFT JOIN routing_sets rs ON rs.id=tr.routing_set_id WHERE tr.task_key=?",(task_key,)).fetchone()
    return {"task":dict(row) if row else {"task_key":task_key,"routing_set_id":None,"routing_set_name":None,"updated_at":None}}

@app.get("/api/routing/tasks")
def list_task_routes(request:Request):
    require_admin(request)
    with closing(get_db()) as db:
        rows=db.execute("SELECT tr.task_key,tr.routing_set_id,rs.name AS routing_set_name,tr.updated_at FROM task_routes tr LEFT JOIN routing_sets rs ON rs.id=tr.routing_set_id ORDER BY tr.task_key").fetchall()
    return {"tasks":[dict(x) for x in rows]}

@app.post("/api/files")
async def upload_file(request:Request):
    user=current_user(request)
    form=await request.form()
    upload=form.get("file")
    if upload is None or not getattr(upload,"filename",None): raise HTTPException(400,"Файл не выбран")
    filename=Path(str(upload.filename)).name
    if _file_extension(filename) not in ALLOWED_FILE_EXTENSIONS: raise HTTPException(400,"Этот тип файла не поддерживается")
    data=await upload.read()
    if len(data)>MAX_FILE_SIZE: raise HTTPException(413,f"Файл слишком большой. Максимум: {MAX_FILE_SIZE // 1024 // 1024} МБ")
    _validate_archive_payload(filename,data)
    target=(UPLOAD_DIR/_safe_stored_filename(filename)).resolve()
    try:
        target.relative_to(UPLOAD_DIR.resolve())
    except ValueError:
        raise HTTPException(400,"Некорректный путь файла")
    target.write_bytes(data)
    mime=upload.content_type or mimetypes.guess_type(filename)[0] or "application/octet-stream"
    try:
        with closing(get_db()) as db:
            cur=db.execute("INSERT INTO files(user_id,filename,stored_filename,mime_type,size,path,created_at) VALUES(?,?,?,?,?,?,?)",(user["id"],filename,target.name,mime,len(data),str(target),now_iso()))
            db.commit()
            row=db.execute("SELECT id,filename,mime_type,size,created_at FROM files WHERE id=?",(cur.lastrowid,)).fetchone()
        return {"file":dict(row)}
    except Exception:
        target.unlink(missing_ok=True)
        raise

@app.get("/api/files/{file_id}/download")
def download_file(file_id:int,request:Request):
    user=current_user(request)
    with closing(get_db()) as db:
        row=db.execute("SELECT path,filename,mime_type FROM files WHERE id=? AND user_id=?",(file_id,user["id"])).fetchone()
    if row is None:
        raise HTTPException(404,"Файл не найден")
    path=_stored_file_path(row["path"])
    if not path.is_file():
        raise HTTPException(404,"Файл не найден")
    return FileResponse(str(path),media_type=row["mime_type"] or "application/octet-stream",filename=row["filename"])

@app.delete("/api/files/{file_id}")
def delete_file(file_id:int,request:Request):
    user=current_user(request)
    with closing(get_db()) as db:
        row=db.execute("SELECT path FROM files WHERE id=? AND user_id=?",(file_id,user["id"])).fetchone()
        if row is None: raise HTTPException(404,"Файл не найден")
        attached=db.execute("SELECT COUNT(*) AS count FROM message_files mf JOIN messages m ON m.id=mf.message_id JOIN chats c ON c.id=m.chat_id WHERE mf.file_id=? AND c.user_id=?",(file_id,user["id"])).fetchone()
        if int(attached["count"])>0:
            raise HTTPException(409,"Файл используется в истории чатов. Сначала удалите сообщения с этим вложением или оставьте файл в библиотеке.")
        db.execute("DELETE FROM files WHERE id=? AND user_id=?",(file_id,user["id"]))
        db.commit()
    path=_stored_file_path(row["path"])
    path.unlink(missing_ok=True)
    return {"ok":True}

@app.get("/api/files")
def list_files(request:Request,q:str=""):
    user=current_user(request)
    query=q.strip()[:MAX_SEARCH_LENGTH]
    with closing(get_db()) as db:
        if query:
            rows=db.execute("SELECT id,filename,mime_type,size,created_at FROM files WHERE user_id=? AND filename LIKE ? ORDER BY created_at DESC",(user["id"],"%"+query+"%")).fetchall()
        else:
            rows=db.execute("SELECT id,filename,mime_type,size,created_at FROM files WHERE user_id=? ORDER BY created_at DESC",(user["id"],)).fetchall()
    return {"files":[{**dict(x),"download_url":f"/api/files/{x['id']}/download"} for x in rows]}

@app.get("/share/{token}")
def shared_frontend(token:str):
    with closing(get_db()) as db:
        row=db.execute("SELECT id FROM chat_shares WHERE token=? AND enabled=1",(token,)).fetchone()
    if row is None:
        raise HTTPException(404,"Ссылка недействительна или отключена")
    index=FRONTEND_DIR/"index.html"
    if not index.exists(): raise HTTPException(404,"Frontend not found")
    return FileResponse(index)

@app.get("/")
def frontend_root():
    index=FRONTEND_DIR/"index.html"
    if not index.exists():
        raise HTTPException(404,"Frontend not found")
    return FileResponse(index)

@app.get("/{path:path}")
def frontend_files(path:str):
    candidate=FRONTEND_DIR/path
    if candidate.is_file():
        return FileResponse(candidate)
    index=FRONTEND_DIR/"index.html"
    if index.exists():
        return FileResponse(index)
    raise HTTPException(404,"Not found")
